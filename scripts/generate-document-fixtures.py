"""Generate only original Vesperwind document test data (requires python-docx and Pillow)."""
from pathlib import Path
from io import BytesIO
from zipfile import ZipFile, ZIP_DEFLATED
from xml.etree import ElementTree as ET
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw

DEST = Path(__file__).resolve().parents[1] / 'test' / 'fixtures' / 'document'
DEST.mkdir(parents=True, exist_ok=True)

def save(name, doc):
    doc.save(DEST / name)

doc = Document()
doc.add_heading('Vesperwind simple document', 0)
doc.add_paragraph('The quick brown fox jumps over the lazy dog.')
doc.add_paragraph('Edit this known paragraph and preserve all other content.')
save('simple.docx', doc)

doc = Document()
doc.add_heading('Formatting sample', 1)
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.left_indent = Inches(.4)
p.paragraph_format.line_spacing = 1.5
p.add_run('Bold ').bold = True
p.add_run('Italic ').italic = True
p.add_run('Underline').underline = True
r = p.add_run(' Colored')
r.font.name = 'Arial'
r.font.size = Pt(16)
r.font.color.rgb = RGBColor(0x22, 0x44, 0x88)
doc.add_paragraph('Bullet one', style='List Bullet')
doc.add_paragraph('Number one', style='List Number')
save('formatting.docx', doc)

doc = Document()
doc.add_heading('Tables', 1)
table = doc.add_table(rows=3, cols=3)
for row in range(3):
    for col in range(3): table.cell(row, col).text = f'R{row+1}C{col+1}'
table.cell(1, 0).merge(table.cell(1, 1)).text = 'Merged cells'
save('tables.docx', doc)

image = Image.new('RGB', (320, 160), '#e6eef9')
draw = ImageDraw.Draw(image)
draw.rectangle((25, 25, 295, 135), fill='#1e5ba8')
draw.text((75, 72), 'Vesperwind', fill='white')
png = BytesIO()
image.save(png, format='PNG')
png.seek(0)
doc = Document()
doc.add_paragraph('Original embedded image below:')
doc.add_picture(png, width=Inches(3))
save('images.docx', doc)

doc = Document()
doc.add_paragraph('First page body')
section = doc.sections[0]
section.header.paragraphs[0].text = 'Vesperwind header'
section.footer.paragraphs[0].text = 'Vesperwind footer'
doc.add_page_break()
doc.add_paragraph('Second page body')
save('headers-footers.docx', doc)

doc = Document()
doc.add_heading('Unicode and Cyrillic — 文書', 1)
doc.add_paragraph('Привет, Иван! Ελληνικά العربية 😀')
save('unicode.docx', doc)

doc = Document()
doc.add_paragraph('Advanced marker: preserve this surrounding paragraph.')
p = doc.add_paragraph('Bookmark, field, and opaque custom XML follow.')
bookmark_start = OxmlElement('w:bookmarkStart')
bookmark_start.set(qn('w:id'), '1')
bookmark_start.set(qn('w:name'), 'VesperwindBookmark')
p._p.insert(0, bookmark_start)
bookmark_end = OxmlElement('w:bookmarkEnd')
bookmark_end.set(qn('w:id'), '1')
p._p.append(bookmark_end)
field = OxmlElement('w:fldSimple')
field.set(qn('w:instr'), ' DATE \\@ "yyyy-MM-dd" ')
run = OxmlElement('w:r')
text = OxmlElement('w:t')
text.text = '2026-09-28'
run.append(text)
field.append(run)
p._p.append(field)
save('advanced.docx', doc)

advanced = DEST / 'advanced.docx'
with ZipFile(advanced) as archive:
    members = {name: archive.read(name) for name in archive.namelist()}
content_types = ET.fromstring(members['[Content_Types].xml'])
ET.SubElement(content_types, '{http://schemas.openxmlformats.org/package/2006/content-types}Override',
              PartName='/customXml/item1.xml', ContentType='application/xml')
members['[Content_Types].xml'] = ET.tostring(content_types, encoding='utf-8', xml_declaration=True)
relationships = ET.fromstring(members['word/_rels/document.xml.rels'])
ET.SubElement(relationships, '{http://schemas.openxmlformats.org/package/2006/relationships}Relationship',
              Id='rIdVesperwindCustom', Type='http://schemas.openxmlformats.org/officeDocument/2006/relationships/customXml', Target='../customXml/item1.xml')
members['word/_rels/document.xml.rels'] = ET.tostring(relationships, encoding='utf-8', xml_declaration=True)
members['customXml/item1.xml'] = b'<?xml version="1.0" encoding="UTF-8"?><vesperwind xmlns="urn:vesperwind:test"><opaque>DO_NOT_REMOVE_12345</opaque></vesperwind>'
with ZipFile(advanced, 'w', ZIP_DEFLATED) as archive:
    for name, data in members.items(): archive.writestr(name, data)

(DEST / 'simple.rtf').write_bytes(b'{\\rtf1\\ansi\\deff0{\\fonttbl{\\f0 Arial;}}\\f0\\fs24 Vesperwind simple RTF\\par Plain paragraph.}')
(DEST / 'formatting.rtf').write_bytes(b'{\\rtf1\\ansi\\deff0{\\fonttbl{\\f0 Arial;}{\\f1 Times New Roman;}}{\\colortbl;\\red34\\green68\\blue136;}\\f0\\fs28\\qc Centered\\par \\ql\\b Bold\\b0  \\i Italic\\i0  \\ul Underline\\ul0  \\cf1 Colored\\par \\f1 Second font.}')
(DEST / 'unicode.rtf').write_bytes(b'{\\rtf1\\ansi\\uc1\\deff0{\\fonttbl{\\f0 Arial;}}\\f0\\fs24 \\u1055?\\u1088?\\u1080?\\u1074?\\u1077?\\u1090?, Vesperwind! \\u25991?\\u26723?\\par}')

long_doc = Document()
for page in range(50):
    long_doc.add_heading(f'Page {page+1}', 1)
    for line in range(8): long_doc.add_paragraph(f'Vesperwind generated page {page+1}, paragraph {line+1}. ' * 5)
    if page < 49: long_doc.add_page_break()
save('long-50-pages.docx', long_doc)
